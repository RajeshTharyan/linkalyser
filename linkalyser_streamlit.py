import asyncio
import aiohttp
from bs4 import BeautifulSoup
import PyPDF2
import docx
import openpyxl
from urllib.parse import urljoin
from io import BytesIO
import streamlit as st
import webbrowser
import requests

# Global stop flag
stop_search = False

def parse_pdf(content):
    try:
        reader = PyPDF2.PdfReader(BytesIO(content))
        return "\f".join(p.extract_text() or "" for p in reader.pages)
    except:
        return ""

def parse_word(content):
    try:
        doc = docx.Document(BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs)
    except:
        return ""

def parse_excel(content):
    try:
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                lines.append(" ".join(str(c) for c in row if c is not None))
        return "\n".join(lines)
    except:
        return ""

def parse_html(content):
    return BeautifulSoup(content, 'html.parser').get_text() or ""

def search_keywords(text, keywords):
    pages = text.split('\f')
    found = {}
    for i, page in enumerate(pages, start=1):
        low = page.lower()
        for kw in keywords:
            if kw.lower() in low:
                found.setdefault(kw, []).append(i)
    return found

async def fetch_and_parse(session, url, stats):
    try:
        async with session.get(url) as resp:
            resp.raise_for_status()
            ct = resp.headers.get('content-type','')
            data = await resp.read()
            if 'application/pdf' in ct:
                stats['PDF'] += 1
                return url, parse_pdf(data)
            elif 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' in ct:
                stats['Word Document'] += 1
                return url, parse_word(data)
            elif 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in ct:
                stats['Excel'] += 1
                return url, parse_excel(data)
            elif 'text/html' in ct:
                stats['HTML'] += 1
                return url, parse_html(data)
            elif 'image/' in ct:
                stats['Image'] += 1
                return url, ""
            else:
                stats['Other'] += 1
                return url, ""
    except:
        return url, ""

async def analyse_links_async(base_url, links, stats, progress_callback):
    results = []
    timeout = aiohttp.ClientTimeout(total=30)
    sem = asyncio.Semaphore(10)  # limit concurrency
    async with aiohttp.ClientSession(timeout=timeout) as session:
        tasks = []
        for link in links:
            async def bound_fetch(link=link):
                async with sem:
                    return await fetch_and_parse(session, link, stats)
            tasks.append(bound_fetch())
        for i, coro in enumerate(asyncio.as_completed(tasks), start=1):
            if stop_search:
                break
            url, content = await coro
            results.append((url, content))
            progress_callback(i, len(links))
    return results

async def keyword_search_async(link_contents, keywords, progress_callback):
    results = []
    for i, (url, content) in enumerate(link_contents, start=1):
        if stop_search:
            break
        found = search_keywords(content, keywords)
        if found:
            results.append({'url': url, 'keywords': found})
        progress_callback(i, len(link_contents))
    return results

def display_statistics(stats, header="Linkalyser Statistics"):
    st.write(f"### {header}")
    st.write(" | ".join(f"{k}: {v}" for k,v in stats.items()))

def run_analysis(url, keywords):
    global stop_search
    stop_search = False

    # initialize stats
    stats = {
        'HTML':0, 'PDF':0, 'Word Document':0,
        'Excel':0, 'Image':0, 'Other':0
    }

    # fetch start page & extract links
    html = requests.get(url).text
    soup = BeautifulSoup(html, 'html.parser')
    links = [
        urljoin(url, a['href'])
        for a in soup.find_all('a', href=True)
        if not a['href'].startswith('mailto:')
    ]

    st.write(f"Found {len(links)} link{'s' if len(links)!=1 else ''} to analyse.")

    # Phase 1: async fetch & parse
    phase1_bar = st.progress(0)
    def p1(i, total): phase1_bar.progress(i/total)
    link_contents = asyncio.run(analyse_links_async(url, links, stats, p1))

    display_statistics(stats, header="After Content Analysis")

    # Phase 2: keyword search
    phase2_bar = st.progress(0)
    def p2(i, total): phase2_bar.progress(i/total)
    results = asyncio.run(keyword_search_async(link_contents, keywords, p2))

    display_statistics(stats, header="Final Statistics")
    st.write("### Keyword Search Results")
    if not results:
        st.write("No Keyword(s) found!")
    for r in results:
        kw_str = ", ".join(f"{k} (Page: {','.join(map(str,v))})"
                            for k,v in r['keywords'].items())
        st.markdown(f"- **URL:** [{r['url']}]({r['url']})  \n  **Keywords:** {kw_str}")

# --- Streamlit UI ---
st.title("🔍 Linkalyser (asyncio)")

url = st.text_input("Enter the URL to start the search:")
kw_input = st.text_input("Enter keywords (comma‑separated):")

c1, c2, c3 = st.columns(3)
if c1.button("Submit"):
    if not url or not kw_input:
        st.error("Please provide both URL and keywords.")
    else:
        run_analysis(url, [k.strip() for k in kw_input.split(",")])
if c2.button("Stop"):
    stop_search = True
if c3.button("Reset"):
    stop_search = False
    st.experimental_rerun()
